"""Render model-authored text into real, valid office documents.

One generation path shared by `create_file` (download links) and
`email_with_attachment` (email attachments), so a `.pdf`/`.docx`/`.xlsx`
filename always yields bytes that actually open in a viewer — never plain
text mislabeled with an office MIME type.

All generators are pure-Python (fpdf2 / python-docx / openpyxl), so they run
on serverless with no system/native dependencies. Input is plain text: DOCX
splits it into paragraphs, XLSX parses it as CSV rows, PDF lays it out
line-by-line. Markdown syntax is not interpreted — it's shown literally.
"""
from __future__ import annotations

import csv
import io
import mimetypes

# Office Open XML MIME types (mimetypes.guess_type is unreliable for these
# across platforms, so pin them explicitly).
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"


def render_pdf(content: str) -> bytes:
    """Render plain text into a REAL, valid PDF. fpdf2's built-in core font is
    Latin-1 only, so common "smart" punctuation is normalized to ASCII first;
    anything else outside Latin-1 is replaced rather than crashing the render.
    """
    from fpdf import FPDF

    normalized = (
        content.replace("‘", "'").replace("’", "'")
        .replace("“", '"').replace("”", '"')
        .replace("–", "-").replace("—", "-")
        .replace("…", "...")
    )
    safe = normalized.encode("latin-1", errors="replace").decode("latin-1")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in safe.split("\n"):
        # multi_cell(w=0, ...) leaves the cursor at the right edge afterward
        # rather than resetting to the left margin — without this, the next
        # call computes an ~empty width and raises FPDFException.
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 6, line or " ")
    return bytes(pdf.output())


def render_docx(content: str) -> bytes:
    """Render plain text into a Word .docx, one paragraph per line."""
    from docx import Document

    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_xlsx(content: str) -> bytes:
    """Render CSV-shaped text into an Excel .xlsx (one worksheet). Content is
    parsed with the csv module, so quoted fields and embedded commas are handled
    correctly; a plain single-column list works too (each line -> one cell)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in csv.reader(io.StringIO(content)):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_document(
    filename: str, content: str, content_type: str | None = None
) -> tuple[bytes, str]:
    """Turn text `content` into file bytes based on `filename`'s extension.

    Returns ``(raw_bytes, content_type)``. `.pdf`/`.docx`/`.xlsx` are rendered
    into real office files; anything else is written as UTF-8 text with a
    guessed (or caller-supplied) MIME type.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        return render_pdf(content), PDF_MIME
    if ext == "docx":
        return render_docx(content), DOCX_MIME
    if ext == "xlsx":
        return render_xlsx(content), XLSX_MIME
    ct = (content_type or "").strip() or mimetypes.guess_type(filename)[0] or "text/plain"
    return content.encode("utf-8"), ct
